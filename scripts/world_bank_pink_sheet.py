"""Verified multi-category physical commodity price benchmarks from the World Bank.

The Pink Sheet's visible "Monthly Prices" worksheet is a broad, non-index
commodity catalog. Its first 4 rows are title/metadata; row 5 contains exact
benchmark names, row 6 units, and row 7+ monthly observations.
Do NOT use the workbook's hidden first worksheet or index worksheet.
"""
from __future__ import annotations

from datetime import date, datetime
from html.parser import HTMLParser
from io import BytesIO
import math
import re
from urllib.parse import urljoin, urlparse

import openpyxl
import requests

LANDING = "https://www.worldbank.org/en/research/commodity-markets"
# Fallback to the current verified official URL if the landing-page structure changes.
FALLBACK = ("https://thedocs.worldbank.org/en/doc/"
            "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/"
            "CMO-Historical-Data-Monthly.xlsx")
CATEGORIES = (
    "Energy", "Precious metals", "Industrial metals", "Forest products",
    "Grains", "Oilseeds & oils", "Soft commodities", "Livestock & food",
    "Fertilizers", "Other commodities",
)
UNITS = {
    "($/bbl)": "USD/barrel",
    "($/mt)": "USD/metric ton",
    "($/kg)": "USD/kg",
    "($/mmbtu)": "USD/MMBtu",
    "($/cubic meter)": "USD/m³",
    "($/troy oz)": "USD/troy oz",
    "($/dmtu)": "USD/dmtu",
    "(cents/sheet)": "cents/sheet",
}


class WorkbookLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            href = dict(attrs).get("href") or ""
            if "CMO-Historical-Data-Monthly.xlsx" in href:
                self.links.append(href)


def select_workbook_url(html: str) -> str:
    parser = WorkbookLinks()
    parser.feed(html)
    for link in parser.links:
        resolved = urljoin(LANDING, link)
        parsed = urlparse(resolved)
        if (parsed.scheme == "https"
                and parsed.netloc == "thedocs.worldbank.org"
                and parsed.path.endswith("/CMO-Historical-Data-Monthly.xlsx")):
            return resolved
    return FALLBACK


def fetch_workbook() -> tuple[bytes, str]:
    session = requests.Session()
    session.headers.update({"User-Agent": "FINC-560 educational commodity dashboard",
                            "Accept": "text/html,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
    try:
        landing = session.get(LANDING, timeout=25)
        landing.raise_for_status()
        url = select_workbook_url(landing.text)
    except requests.RequestException:
        url = FALLBACK
    response = session.get(url, timeout=55)
    response.raise_for_status()
    if len(response.content) < 25_000 or len(response.content) > 12_000_000:
        raise ValueError("Invalid World Bank workbook response size")
    if not response.content.startswith(b"PK"):
        raise ValueError("Official World Bank file is not a valid Excel workbook")
    return response.content, url


def category_for(name: str) -> str:
    n = name.lower()
    if any(x in n for x in ("crude oil", "natural gas", "liquefied natural gas", "coal,")):
        return "Energy"
    if n in ("gold", "platinum", "silver"):
        return "Precious metals"
    if any(x in n for x in ("aluminum", "iron ore", "copper", "lead", "tin",
                             "nickel", "zinc", "uranium", "cobalt", "lithium")):
        return "Industrial metals"
    if any(x in n for x in ("logs,", "sawnwood,", "plywood")):
        return "Forest products"
    if any(x in n for x in ("barley", "maize", "sorghum", "rice,", "wheat,")):
        return "Grains"
    if any(x in n for x in ("soybean", "coconut oil", "groundnut", "palm",
                             "rapeseed", "sunflower")):
        return "Oilseeds & oils"
    if any(x in n for x in ("cocoa", "coffee", "tea,", "sugar", "tobacco", "cotton",
                             "rubber", "wool")):
        return "Soft commodities"
    if any(x in n for x in ("beef", "chicken", "lamb", "shrimp",
                             "fish meal", "banana", "orange", "milk")):
        return "Livestock & food"
    if any(x in n for x in ("phosphate", "potassium chloride", "urea")) or n in ("dap", "tsp"):
        return "Fertilizers"
    return "Other commodities"


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def next_month(iso: str) -> str:
    year, month = map(int, iso[:7].split("-"))
    return f"{year+1 if month==12 else year:04d}-{1 if month==12 else month+1:02d}-01"


def parse_monthly_workbook(payload: bytes, url: str, today: datetime) -> tuple[list[dict], list[dict]]:
    workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if "Monthly Prices" not in workbook.sheetnames:
            raise ValueError("World Bank Monthly Prices worksheet missing")
        sheet = workbook["Monthly Prices"]
        rows = sheet.iter_rows(values_only=True)
        # Explicitly select the real prices worksheet, never the hidden metadata tab.
        for _ in range(4):
            next(rows)
        names = next(rows)
        units = next(rows)
        if "Crude oil, Brent" not in names or "Gold" not in names:
            raise ValueError("World Bank monthly price column schema has changed")
        current_month = today.date().replace(day=1)
        columns = {}
        for col in range(1, min(len(names), len(units))):
            raw_name, raw_unit = names[col], units[col]
            if not isinstance(raw_name, str) or not isinstance(raw_unit, str):
                continue
            name = raw_name.replace("**", "").strip()
            unit = UNITS.get(raw_unit.strip())
            # Never treat a commodity INDEX (2010=100) as a physical price.
            if not name or not unit:
                continue
            key = "wb-" + slugify(name)
            if key in columns:
                raise ValueError(f"Duplicate World Bank benchmark {key}")
            columns[key] = (col, name, unit, [])
        if len(columns) < 40:
            raise ValueError("World Bank workbook unexpectedly lacks physical benchmarks")
        for row in rows:
            period = row[0] if row else None
            if not isinstance(period, str) or not re.fullmatch(r"\d{4}M(0[1-9]|1[0-2])", period):
                continue
            current = date(int(period[:4]), int(period[5:]), 1)
            if current >= current_month:
                continue  # Incomplete current month must never look finalized.
            stamp = current.isoformat()
            for col, _name, _unit, observations in columns.values():
                v = row[col] if col < len(row) else None
                if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and v > 0:
                    observations.append({"date": stamp, "value": round(float(v), 6)})

        selected, skipped = [], []
        for key, (_col, name, unit, observations) in columns.items():
            if not observations:
                skipped.append({"name": name, "reason": "no verified positive observations"})
                continue
            recent = observations[-1]["date"]
            age_days = (current_month - date.fromisoformat(recent)).days
            # A regression treating nonadjacent calendar months as equally spaced
            # misrepresents uncertainty. Require 60 contiguous recent observations.
            trailing = observations[-60:]
            contiguous = len(trailing) >= 60 and all(
                next_month(trailing[i-1]["date"]) == trailing[i]["date"]
                for i in range(1, len(trailing))
            )
            if age_days > 125 or not contiguous:
                skipped.append({
                    "name": name,
                    "reason": "stale series" if age_days > 125 else
                              "fewer than 60 consecutive recent monthly prices",
                    "last_observation": recent,
                })
                continue
            selected.append({
                "id": key,
                "label": name,
                "category": category_for(name),
                "unit": unit,
                "source": "World Bank Commodity Price Data (Pink Sheet)",
                "source_id": "WB:" + slugify(name),
                "source_url": url,
                "frequency": "published monthly benchmark price (not exchange futures)",
                "last_observation": recent,
                "observations": observations,
            })
        if len(selected) < 30:
            raise ValueError(f"World Bank coverage unexpectedly narrow: {len(selected)} physical prices")
        return selected, skipped
    finally:
        workbook.close()
